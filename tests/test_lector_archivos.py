"""Tests del lector de archivos (B3a). Ejecutar: python -m unittest discover tests"""
import io
import sys
import unittest
from pathlib import Path
from openpyxl import Workbook

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))
from lector_archivos import ErrorLectura, leer


def xlsx(*hojas, combinar=None):
    wb = Workbook()
    wb.remove(wb.active)
    for titulo, filas in hojas:
        ws = wb.create_sheet(titulo)
        for f in filas:
            ws.append(f)
        for rango in combinar or []:
            ws.merge_cells(rango)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestLector(unittest.TestCase):
    def test_csv_separadores_y_codificaciones(self):
        for sep in [";", ",", "\t"]:
            texto = sep.join(["Nombre", "Cuota", "Alta"]) + "\r\n" + sep.join(["Ana Núñez", "35,50" if sep != "," else "35.50", "01/02/2026"]) + "\r\n"
            for cod in ["utf-8", "utf-8-sig", "cp1252"]:
                t, _ = leer("socios.csv", texto.encode(cod))
                self.assertEqual(t[0].cabecera, ["Nombre", "Cuota", "Alta"], (sep, cod))
                self.assertEqual(t[0].filas[0][0], "Ana Núñez", (sep, cod))

    def test_titulo_filas_vacias_total_y_duplicado(self):
        datos = xlsx(("Hoja", [["CLUB INVENTADO"], [], ["Nombre", "Telefono", "Cuota"], ["Ana Gil", "600000001", 30],
                              [None, None, None], ["  Luis Paz ", "600000002", 40], ["Luis Paz", "600000002", 40],
                              ["Nombre", "Telefono", "Cuota"], ["TOTAL", None, 70]]))
        t, _ = leer("a.xlsx", datos)
        self.assertEqual(t[0].fila_cabecera, 3)
        self.assertEqual([f[0] for f in t[0].filas], ["Ana Gil", "Luis Paz"])
        self.assertEqual(t[0].num_filas, [4, 6])
        self.assertEqual(len(t[0].avisos), 4)                   # titulo, duplicado, cabecera repetida, total

    def test_cabecera_doble_combinada(self):
        datos = xlsx(("Hoja", [["SOCIO", None, "PAGOS"], ["Nombre", "Movil", "Octubre"], ["Ana Gil", "600000001", "Si"]]),
                     combinar=["A1:B1"])
        t, _ = leer("a.xlsx", datos)
        self.assertEqual(t[0].cabecera, ["SOCIO / Nombre", "SOCIO / Movil", "PAGOS / Octubre"])
        self.assertEqual(len(t[0].filas), 1)

    def test_varias_hojas_y_hoja_de_notas(self):
        datos = xlsx(("Altas", [["Nombre", "Cuota"], ["Ana Gil", 30]]), ("Bajas", [["Nombre", "Cuota"], ["Luis Paz", 40]]),
                     ("Notas", [["Revisar cuotas en enero"]]))
        t, avisos = leer("a.xlsx", datos)
        self.assertEqual([x.hoja for x in t], ["Altas", "Bajas"])
        self.assertEqual(len(avisos), 1)

    def test_errores_claros(self):
        for nombre, datos in [("a.xls", b"x"), ("a.pdf", b"x"), ("a.xlsx", b"no es un excel"), ("a.csv", b"")]:
            with self.assertRaises(ErrorLectura):
                leer(nombre, datos)

    def test_examen_listas_sin_excepcion(self):
        """Hecho cuando de B3: los 60 archivos de examen/listas/ se leen sin excepcion (solo se mide, no se ajusta)."""
        archivos = sorted((RAIZ / "examen" / "listas").glob("*"))
        self.assertEqual(len(archivos), 60)
        for a in archivos:
            with self.subTest(archivo=a.name):
                tablas, _ = leer(a.name, a.read_bytes())
                self.assertTrue(tablas and all(t.filas for t in tablas))


if __name__ == "__main__":
    unittest.main()
