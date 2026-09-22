import re
from datetime import date, datetime, time
import pandas as pd
import streamlit as st
from openpyxl import load_workbook

st.set_page_config(page_title="Gestion del club", layout="wide")
CUOTA, PLAZAS = 65, 14
DIAS = ["LUNES", "MARTES", "MIERCOLES", "JUEVES", "VIERNES", "SABADO", "DOMINGO"]
MESES = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO",
         "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]
TILDES = str.maketrans("ÁÉÍÓÚáéíóú", "AEIOUaeiou")
RE_DIA = re.compile(r"^(LUNES|MARTES|MIERCOLES|JUEVES|VIERNES|SABADO|DOMINGO)\s+(\d{1,2})")
RE_CP = re.compile(r"\s*\bCP\s*\d*\s*$", re.I)


def limpiar(texto):
    return str(texto).translate(TILDES).upper().strip()


def elegir_anio(titulo, mes, cabeceras, anterior):
    m = re.search(r"(20\d{2}|\b\d{2}\b)", titulo)
    if m:
        a = int(m.group(1))
        return a + 2000 if a < 100 else a
    dia_sem, dia = cabeceras[0]
    base = anterior or date.today().year
    for a in sorted(range(base - 3, base + 4), key=lambda x: (x < base, abs(x - base))):
        try:
            if date(a, mes, dia).weekday() == dia_sem:
                return a
        except ValueError:
            pass
    return base


@st.cache_data
def leer_excel(fuente):
    wb = load_workbook(fuente, data_only=True, read_only=True)
    filas_out, anio = [], None
    for ws in wb.worksheets:
        titulo = limpiar(ws.title)
        mes = next((i + 1 for i, n in enumerate(MESES) if re.search(rf"\b{n}\b", titulo)), None)
        if not mes:
            continue
        filas = list(ws.iter_rows(values_only=True))
        bloques = []
        for i, f in enumerate(filas):
            if f and isinstance(f[0], str):
                m = RE_DIA.match(limpiar(f[0]))
                if m:
                    bloques.append((i, DIAS.index(m.group(1)), int(m.group(2))))
        if not bloques:
            continue
        anio = elegir_anio(titulo, mes, [(b[1], b[2]) for b in bloques], anio)
        for k, (i, _, dia) in enumerate(bloques):
            fin = bloques[k + 1][0] if k + 1 < len(bloques) else len(filas)
            horas = filas[i + 2] if i + 2 < len(filas) else ()
            for f in filas[i + 3:fin]:
                for c, v in enumerate(f or ()):
                    h = horas[c] if c < len(horas) else None
                    if not isinstance(v, str) or not isinstance(h, (time, datetime)):
                        continue
                    nombre = re.sub(r"^\s*\d+\s*\.\s*", "", v).strip()
                    if not nombre:
                        continue
                    prueba = bool(RE_CP.search(nombre))
                    nombre = RE_CP.sub("", nombre).strip() or nombre
                    try:
                        fecha = date(anio, mes, dia)
                    except ValueError:
                        continue
                    filas_out.append({"fecha": fecha, "hora": h.strftime("%H:%M"),
                                      "nombre": nombre.title(), "prueba": prueba})
    df = pd.DataFrame(filas_out)
    df["fecha"] = pd.to_datetime(df["fecha"])
    df["clave"] = df["nombre"].str.lower().str.replace(r"[.\s]+", " ", regex=True).str.strip()
    return df


def cargar_pagos():
    try:
        p = pd.read_csv("pagos.csv", parse_dates=["fecha_pago"])
    except FileNotFoundError:
        p = pd.DataFrame(columns=["nombre", "mes", "fecha_pago", "importe", "metodo"])
    p["clave"] = p["nombre"].astype(str).str.lower().str.replace(r"[.\s]+", " ", regex=True).str.strip()
    return p


st.sidebar.header("Datos")
subido = st.sidebar.file_uploader("Excel de horarios del club", type="xlsx")
st.sidebar.caption("Sin archivo se usa la demo con datos ficticios. "
                   "No subas datos reales a una app publica.")
res = leer_excel(subido if subido else "horarios_demo.xlsx")
res = res[res["fecha"] <= pd.Timestamp(date.today())]
pagos = cargar_pagos()
if res.empty:
    st.error("No he encontrado reservas en ese Excel")
    st.stop()

hoy = res["fecha"].max()
mes_actual = hoy.strftime("%Y-%m")
st.title("Gestion del club de boxeo")
st.caption(f"Datos hasta el {hoy:%d/%m/%Y} · cuota {CUOTA} € al mes")

socios_res = res[~res["prueba"]]
g = socios_res.groupby("clave")
socios = pd.DataFrame({
    "nombre": g["nombre"].last(),
    "clases este mes": g["fecha"].apply(lambda f: (f.dt.strftime("%Y-%m") == mes_actual).sum()),
    "clases totales": g.size(),
    "ultima clase": g["fecha"].max(),
    "horario habitual": g["hora"].agg(lambda h: h.mode().iat[0]),
})
socios["dias sin venir"] = (hoy - socios["ultima clase"]).dt.days
pagos_mes = pagos[pagos["mes"] == mes_actual].groupby("clave")["fecha_pago"].max()
socios["pagado este mes"] = socios.index.isin(pagos_mes.index)
socios["fecha de pago"] = pagos_mes.reindex(socios.index)
ult_pago = pagos.groupby("clave")["fecha_pago"].max()
socios["ultimo pago"] = ult_pago.reindex(socios.index)
activos = socios[socios["dias sin venir"] <= 30]

def estado(r):
    if r["dias sin venir"] > 30:
        return "Inactivo"
    if not r["pagado este mes"]:
        return "Pendiente de pago"
    if r["dias sin venir"] > 14:
        return "Riesgo de baja"
    return "Al dia"
ORDEN = ["Pendiente de pago", "Riesgo de baja", "Al dia", "Inactivo"]
socios["estado"] = pd.Categorical(socios.apply(estado, axis=1), ORDEN, ordered=True)

pruebas = res[res["prueba"]].groupby("clave")["fecha"].min()
convertidos = [c for c, f in pruebas.items()
               if c in socios.index and socios.loc[c, "ultima clase"] > f]

cobrado = pagos.loc[pagos["mes"] == mes_actual, "importe"].sum()
c1, c2, c3, c4, c5 = st.columns(5)
c1.metric("Socios activos (30 dias)", len(activos))
c2.metric(f"Cobrado {mes_actual}", f"{cobrado:.0f} / {len(activos) * CUOTA} €")
c3.metric("Pendientes de pago", int((socios["estado"] == "Pendiente de pago").sum()))
c4.metric("Riesgo de baja (+14 dias)", int((socios["estado"] == "Riesgo de baja").sum()))
c5.metric("Pruebas que se apuntan", f"{len(convertidos)}/{len(pruebas)}")

t1, t2, t3, t4 = st.tabs(["Socios", "Ficha de socio", "Registrar pago", "Horarios"])

with t1:
    filtro = st.multiselect("Estado", ORDEN, default=ORDEN[:3])
    buscar = st.text_input("Buscar socio")
    vista = socios[socios["estado"].isin(filtro)]
    if buscar:
        vista = vista[vista["nombre"].str.contains(buscar, case=False)]
    st.dataframe(
        vista.sort_values(["estado", "dias sin venir"], ascending=[True, False])
             .assign(**{"fecha de pago": lambda v: v["fecha de pago"].dt.strftime("%d/%m/%Y").fillna("-")}),
        hide_index=True, width="stretch",
        column_order=["nombre", "estado", "pagado este mes", "fecha de pago", "clases este mes",
                      "ultima clase", "dias sin venir", "horario habitual", "clases totales"],
        column_config={
            "ultima clase": st.column_config.DateColumn(format="DD/MM/YYYY"),
            "pagado este mes": st.column_config.CheckboxColumn(),
        })

with t2:
    clave = st.selectbox("Socio", socios.sort_values("nombre").index,
                         format_func=lambda c: socios.loc[c, "nombre"])
    s = socios.loc[clave]
    a, b, c, d = st.columns(4)
    a.metric("Estado", s["estado"])
    b.metric("Dias sin venir", int(s["dias sin venir"]))
    c.metric("Clases este mes", int(s["clases este mes"]))
    d.metric("Ultimo pago", s["ultimo pago"].strftime("%d/%m/%Y") if pd.notna(s["ultimo pago"]) else "Nunca")
    izq, der = st.columns(2)
    izq.subheader("Clases reservadas")
    izq.dataframe(res[res["clave"] == clave].sort_values("fecha", ascending=False)
                  [["fecha", "hora", "prueba"]], hide_index=True, width="stretch",
                  column_config={"fecha": st.column_config.DateColumn(format="DD/MM/YYYY")})
    der.subheader("Pagos")
    der.dataframe(pagos[pagos["clave"] == clave].sort_values("fecha_pago", ascending=False)
                  [["mes", "fecha_pago", "importe", "metodo"]], hide_index=True, width="stretch",
                  column_config={"fecha_pago": st.column_config.DateColumn(format="DD/MM/YYYY")})

with t3:
    st.caption("En la version publica los pagos anadidos se pierden al reiniciarse la app.")
    with st.form("pago", clear_on_submit=True):
        clave_p = st.selectbox("Socio", socios.sort_values("nombre").index,
                               format_func=lambda c: socios.loc[c, "nombre"])
        col1, col2, col3, col4 = st.columns(4)
        mes_p = col1.text_input("Mes (AAAA-MM)", mes_actual)
        fecha_p = col2.date_input("Fecha de pago", date.today(), format="DD/MM/YYYY")
        importe_p = col3.number_input("Importe (€)", min_value=0.0, value=float(CUOTA), step=5.0)
        metodo_p = col4.selectbox("Metodo", ["Bizum", "Efectivo", "Tarjeta", "Transferencia"])
        if st.form_submit_button("Guardar pago"):
            if not re.fullmatch(r"\d{4}-\d{2}", mes_p):
                st.error("El mes debe tener el formato AAAA-MM, por ejemplo 2026-10")
            else:
                nuevo = pd.DataFrame([[socios.loc[clave_p, "nombre"], mes_p, fecha_p, importe_p, metodo_p]],
                                     columns=["nombre", "mes", "fecha_pago", "importe", "metodo"])
                nuevo.to_csv("pagos.csv", mode="a", header=False, index=False)
                st.success(f"Pago guardado: {socios.loc[clave_p, 'nombre']} · {mes_p} · {importe_p:.0f} €")
                st.rerun()

with t4:
    lv = res[res["fecha"].dt.dayofweek < 5]
    clases = lv.groupby(["fecha", "hora"]).size().reset_index(name="n")
    clases["dia"] = clases["fecha"].dt.dayofweek.map(lambda x: DIAS[x].capitalize())
    orden_dias = [x.capitalize() for x in DIAS[:5]]
    tabla = (clases.pivot_table(index="dia", columns="hora", values="n", aggfunc="mean")
             .reindex(orden_dias) / PLAZAS)
    st.subheader("Ocupacion media (lunes a viernes)")
    st.dataframe(tabla.style.format("{:.0%}", na_rep="-")
                 .background_gradient(cmap="RdYlGn", axis=None), width="stretch")
    st.subheader("Apuntados por clase")
    dia_sel = st.date_input("Dia", hoy.date(), format="DD/MM/YYYY")
    del_dia = res[res["fecha"].dt.date == dia_sel]
    if del_dia.empty:
        st.info("No hay reservas ese dia")
    else:
        for hora, grupo in del_dia.groupby("hora"):
            with st.expander(f"{hora} · {len(grupo)}/{PLAZAS} plazas"):
                st.write(", ".join(n + (" (prueba)" if p else "")
                                   for n, p in zip(grupo["nombre"], grupo["prueba"])))
