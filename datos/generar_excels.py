"""Genera Excels inventados tipo horario de club, con su etiqueta correcta por celda (para entrenar B2).
Solo datos inventados. Semilla fija: siempre salen los mismos datos.
Uso: python datos/generar_excels.py  ->  datos/generados/ (excels/*.xlsx, etiquetas.csv, reservas.csv, familias.csv)
familias.csv: el estilo de club (0-29) de cada Excel, para examinar el modelo con formatos que no ha visto (B2b)."""
import calendar
import csv
import random
from datetime import date, datetime, time
from pathlib import Path
from openpyxl import Workbook

N_EXCELS = 300
SEMILLA = 42
SALIDA = Path(__file__).parent / "generados"
# Etiquetas por celda (solo celdas con algo escrito)
DIA, HORA, NOMBRE, PRUEBA, OTRO = "dia", "hora", "nombre", "prueba", "otro"

DIAS = ["LUNES", "MARTES", "MIERCOLES", "JUEVES", "VIERNES", "SABADO", "DOMINGO"]
DIAS_TILDE = ["LUNES", "MARTES", "MIÉRCOLES", "JUEVES", "VIERNES", "SÁBADO", "DOMINGO"]
MESES = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO",
         "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]
NOMBRES = ["Ana", "Luis", "Marta", "Pablo", "Lucia", "Sergio", "Elena", "Jorge", "Sara", "Ivan", "Paula",
           "Raul", "Nuria", "Hugo", "Irene", "Dani", "Carla", "Mario", "Alba", "Ruben", "Noa", "Alex", "Julia",
           "Victor", "Laia", "Adrian", "Clara", "Marcos", "Eva", "Oscar"]
APELLIDOS = ["Garcia", "Lopez", "Martin", "Sanz", "Ruiz", "Moreno", "Navarro", "Gil", "Serrano", "Molina",
             "Castro", "Ortiz", "Rubio", "Marin", "Iglesias", "Vidal", "Prieto", "Soler", "Pastor", "Cano"]
CLASES = ["BOXEO", "FUNCIONAL", "KICK", "BOXEO INICIACION", "CROSS", "PILATES", "MMA", "SPARRING"]
ESTILOS_DIA = ["MAYUS", "Titulo", "tilde", "coma", "corto", "cero", "mes", "barra", "abrev", "dd/mm/aaaa", "d-mes", "fecha"]
MARCAS_PRUEBA = ["CP", "cp", "prueba", "Prueba", "CLASE PRUEBA"]      # lista: columna de observaciones
RUIDO = ["TOTAL", "COMPLETO", "FESTIVO", "X", "OK", "Traer guantes", "SIN CLASE", "cancelada", "ver WhatsApp"]


def errata(r, txt):
    """A veces cambia dos letras de sitio o quita una (erratas tipicas al escribir a mano)."""
    if len(txt) < 5 or r.random() > 0.08:
        return txt
    i = r.randrange(1, len(txt) - 2)
    return txt[:i] + txt[i + 1] + txt[i] + txt[i + 2:] if r.random() < 0.5 else txt[:i] + txt[i + 1:]


def estilo(r):
    """Formato propio de cada 'club': asi los Excels se parecen dentro de un club y varian entre clubs."""
    return {
        "col0": r.choice([1, 1, 1, 2]),                        # columna donde empieza todo
        "dia": r.choice(ESTILOS_DIA),
        "fila_clase": r.random() < 0.6,                       # fila con el nombre de la clase entre dia y horas
        "hora": r.choice(["time", "time", "texto", "texto0", "h", "rango", "datetime", "punto", "mixta"]),
        "numero": r.choice(["1. ", "1.", "1 ", "1) ", "", "1.- "]),
        "cp": r.choice([" CP", " cp", " CP1", " (CP)", " - CP", " CP 1", " (prueba)", " PRUEBA"]),
        "etiqueta_hora": r.choice([None, "HORA", "Horario"]),
        "sabado": r.random() < 0.5,
        "hueco": r.choice([0, 1, 1, 2]),                      # filas vacias entre dias
        "horas": sorted(r.sample([7, 8, 9, 10, 11, 12, 17, 18, 19, 20, 21], r.randint(2, 6))),
        "titulo": r.random() < 0.5,
        # Tipos de formato de B2b-2 (mismo generador, mas variedad)
        "formato": r.choice(["bloques", "bloques", "bloques", "semana", "lista"]),  # ver hoja()
        "cp_delante": r.random() < 0.3,                       # "(CP) Ana Garcia" en vez de "Ana Garcia (CP)"
        "inicial": r.random() < 0.4,                          # "Ana G" en vez de "Ana Garcia"
        "hueco_dia": r.random() < 0.3,                        # fila vacia entre el dia y las horas
        "sep": r.choice([" / ", ", ", "; ", " - "]),         # semana: varios nombres en la misma celda
        "cp_columna": r.random() < 0.6,                       # lista: la marca de prueba va en su propia columna
        "dia_mezcla": r.random() < 0.4,                       # lista: cada fila con un formato de dia distinto
        "desordenada": r.random() < 0.5,                      # lista: filas sin ordenar por fecha
    }


def texto_dia(r, e, anio, mes, dia, forma=None):
    dsem = date(anio, mes, dia).weekday()
    d, m = errata(r, DIAS[dsem]), MESES[mes - 1][:3].lower()
    forma = forma or e["dia"]
    if forma == "fecha":                                         # celda con fecha de verdad (no texto)
        return datetime(anio, mes, dia)
    return {"MAYUS": f"{d} {dia}", "Titulo": f"{d.title()} {dia}", "tilde": f"{DIAS_TILDE[dsem]} {dia}",
            "coma": f"{d.title()}, {dia}", "corto": f"{d[:3]} {dia}", "cero": f"{d} {dia:02d}",
            "mes": f"{d.title()} {dia} {m}", "barra": f"{d.title()} {dia}/{mes:02d}",
            "abrev": f"{DIAS_TILDE[dsem][:r.choice([3, 4])].title()}. {dia}/{mes}",
            "dd/mm/aaaa": f"{dia:02d}/{mes:02d}/{anio}", "d-mes": f"{dia}-{m}"}[forma]


def valor_hora(r, e, h):
    m = r.choice([0, 0, 0, 30])
    forma = r.choice(["time", "texto0"]) if e["hora"] == "mixta" else e["hora"]  # mixta: cada celda a su manera
    return {"punto": f"{h:02d}.{m:02d}h", "time": time(h, m), "texto": f"{h}:{m:02d}", "texto0": f"{h:02d}:{m:02d}", "h": f"{h}h",
            "rango": f"{h}:{m:02d}-{h + 1}:{m:02d}", "datetime": datetime(1900, 1, 1, h, m)}[forma], f"{h:02d}:{m:02d}"


def hoja(r, e, ws, anio, mes, nombre_archivo, etiquetas, reservas):
    """Una hoja con un mes de reservas. Tres tipos de formato (e["formato"]):
    bloques: por cada dia, fila de horas y debajo los nombres en columna.
    semana: tabla por semana, dias en columnas y horas en filas, varios nombres por celda.
    lista: una fila por reserva (fecha | hora | nombre | observaciones)."""
    c0 = e["col0"]

    def poner(f, c, v, etiqueta):
        ws.cell(row=f, column=c, value=v)
        txt = v.date().isoformat() if isinstance(v, datetime) and v.year > 1900 else v.strftime("%H:%M") if hasattr(v, "strftime") else v
        etiquetas.append([nombre_archivo, ws.title, f, c, txt, etiqueta])

    def con_cp(nombre):
        return f"{e['cp'].strip(' -')} {nombre}" if e["cp_delante"] else f"{nombre}{e['cp']}"

    fila = 1
    if e["titulo"]:
        poner(fila, c0, f"HORARIO {r.choice(CLASES)} - {MESES[mes - 1]} {anio}", OTRO)
        fila += 2
    socios = [f"{r.choice(NOMBRES)} {r.choice(APELLIDOS)[:1] if e['inicial'] else r.choice(APELLIDOS)}"
              for _ in range(r.randint(15, 40))]
    ultimo = 6 if e["sabado"] else 5
    dias = [d for d in range(1, calendar.monthrange(anio, mes)[1] + 1) if date(anio, mes, d).weekday() < ultimo]
    {"bloques": bloques, "semana": semana, "lista": lista}[e["formato"]](r, e, anio, mes, nombre_archivo, reservas,
                                                                        poner, con_cp, fila, socios, dias)


def bloques(r, e, anio, mes, nombre_archivo, reservas, poner, con_cp, fila, socios, dias):
    c0 = e["col0"]
    for dia in dias:
        poner(fila, c0, texto_dia(r, e, anio, mes, dia), DIA)
        fila += 1 + e["hueco_dia"]
        if e["fila_clase"]:
            for k in range(len(e["horas"])):
                poner(fila, c0 + 1 + k, r.choice(CLASES), OTRO)
            fila += 1
        if e["etiqueta_hora"]:
            poner(fila, c0, e["etiqueta_hora"], OTRO)
        horas = []
        for k, h in enumerate(e["horas"]):
            v, hhmm = valor_hora(r, e, h)
            poner(fila, c0 + 1 + k, v, HORA)
            horas.append(hhmm)
        fila += 1
        apuntados = [r.sample(socios, min(len(socios), r.choice([0, 2, 4, 6, 8, 10, 14]))) for _ in horas]
        for i in range(max(len(a) for a in apuntados)):
            if r.random() < 0.04:
                fila += 1                                        # fila vacia suelta entre nombres
            for k, lista_k in enumerate(apuntados):
                if i >= len(lista_k):
                    continue
                prueba = r.random() < 0.07
                num = e["numero"].replace("1", str(i + 1))
                poner(fila, c0 + 1 + k, f"{num}{con_cp(lista_k[i]) if prueba else lista_k[i]}", PRUEBA if prueba else NOMBRE)
                reservas.append([nombre_archivo, date(anio, mes, dia).isoformat(), horas[k], lista_k[i], int(prueba)])
            if r.random() < 0.03:
                poner(fila, c0 + 2 + len(horas), r.choice(RUIDO), OTRO)
            fila += 1
        if r.random() < 0.05:
            poner(fila, c0 + 1, r.choice(RUIDO), OTRO)
            fila += 1
        fila += e["hueco"]


def semana(r, e, anio, mes, nombre_archivo, reservas, poner, con_cp, fila, socios, dias):
    c0 = e["col0"]
    semanas = {}
    for dia in dias:
        semanas.setdefault(date(anio, mes, dia).isocalendar()[1], []).append(dia)
    for dias_semana in semanas.values():
        if e["etiqueta_hora"]:
            poner(fila, c0, e["etiqueta_hora"], OTRO)
        for j, dia in enumerate(dias_semana):
            poner(fila, c0 + 1 + j, texto_dia(r, e, anio, mes, dia), DIA)
        fila += 1
        for h in e["horas"]:
            v, hhmm = valor_hora(r, e, h)
            poner(fila, c0, v, HORA)
            for j, dia in enumerate(dias_semana):
                nombres = r.sample(socios, r.choice([0, 0, 1, 1, 2, 3, 4]))
                if not nombres:
                    continue
                pruebas = [r.random() < 0.07 for _ in nombres]
                poner(fila, c0 + 1 + j, e["sep"].join(con_cp(n) if p else n for n, p in zip(nombres, pruebas)),
                      PRUEBA if any(pruebas) else NOMBRE)
                reservas += [[nombre_archivo, date(anio, mes, dia).isoformat(), hhmm, n, int(p)]
                             for n, p in zip(nombres, pruebas)]
            fila += 1
        if r.random() < 0.05:
            poner(fila, c0 + 1, r.choice(RUIDO), OTRO)
            fila += 1
        fila += 1 + e["hueco"]


def lista(r, e, anio, mes, nombre_archivo, reservas, poner, con_cp, fila, socios, dias):
    c0 = e["col0"]
    filas = [(dia, h, n, r.random() < 0.1) for dia in dias for h in e["horas"]
             for n in r.sample(socios, r.choice([0, 0, 1, 2, 3]))]
    if e["desordenada"]:
        r.shuffle(filas)
    obs = e["cp_columna"] or r.random() < 0.5
    cabecera = [r.choice(["Fecha", "FECHA", "Dia"]), r.choice(["Hora", "HORA", "Horario"]),
                r.choice(["Nombre", "NOMBRE", "Socio", "Cliente", "Alumno"])] + ([r.choice(["Observaciones", "Notas", "Obs."])] if obs else [])
    for k, v in enumerate(cabecera):
        poner(fila, c0 + k, v, OTRO)
    fila += 1
    for dia, h, n, prueba in filas:
        v, hhmm = valor_hora(r, e, h)
        poner(fila, c0, texto_dia(r, e, anio, mes, dia, r.choice(ESTILOS_DIA) if e["dia_mezcla"] else None), DIA)
        poner(fila, c0 + 1, v, HORA)
        poner(fila, c0 + 2, con_cp(n) if prueba and not e["cp_columna"] else n, PRUEBA if prueba else NOMBRE)
        if obs and prueba and e["cp_columna"]:
            poner(fila, c0 + 3, r.choice(MARCAS_PRUEBA), OTRO)
        elif obs and r.random() < 0.25:
            poner(fila, c0 + 3, r.choice(RUIDO), OTRO)
        reservas.append([nombre_archivo, date(anio, mes, dia).isoformat(), hhmm, n, int(prueba)])
        fila += 1


def main():
    r = random.Random(SEMILLA)
    (SALIDA / "excels").mkdir(parents=True, exist_ok=True)
    etiquetas, reservas, familias = [], [], []
    clubs = [estilo(r) for _ in range(30)]
    for n in range(N_EXCELS):
        base = r.choice(clubs)
        e = dict(base)
        if r.random() < 0.2:                                     # a veces el club cambia algo de su formato
            e.update({k: v for k, v in estilo(r).items() if r.random() < 0.3})
        anio, mes = r.choice([2024, 2025, 2026]), r.randint(1, 12)
        nombre_archivo = f"horario_{n:03d}.xlsx"
        familias.append([nombre_archivo, next(i for i, c in enumerate(clubs) if c is base)])
        wb = Workbook()
        wb.properties.created = datetime(2026, 1, 1)
        wb.remove(wb.active)
        for m in range(r.choice([1, 1, 2])):
            mm = (mes + m - 1) % 12 + 1
            aa = anio + (mes + m - 1) // 12
            titulo = r.choice([f"{MESES[mm - 1]}", f"{MESES[mm - 1]} {aa}", f"{MESES[mm - 1].title()} {aa % 100}"])
            hoja(r, e, wb.create_sheet(titulo), aa, mm, nombre_archivo, etiquetas, reservas)
        if r.random() < 0.15:
            wb.create_sheet(r.choice(["Hoja1", "RESUMEN", "Notas"]))
        wb.save(SALIDA / "excels" / nombre_archivo)
    for archivo, cab, filas in [("etiquetas.csv", ["archivo", "hoja", "fila", "col", "valor", "etiqueta"], etiquetas),
                                ("reservas.csv", ["archivo", "fecha", "hora", "nombre", "prueba"], reservas),
                                ("familias.csv", ["archivo", "familia"], familias)]:
        with open(SALIDA / archivo, "w", newline="", encoding="utf-8") as f:
            csv.writer(f).writerows([cab] + filas)
    print(f"{N_EXCELS} Excels, {len(etiquetas)} celdas etiquetadas, {len(reservas)} reservas -> {SALIDA}")


if __name__ == "__main__":
    main()
