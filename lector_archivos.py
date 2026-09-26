"""Lectura robusta de listas en Excel o CSV (B3a). Solo encuentra la tabla: no decide que es cada columna (eso es B4).
Uso: tablas, avisos = leer(nombre_archivo, contenido_en_bytes)
Cada Tabla trae la cabecera real (saltando titulos, filas vacias y cabeceras dobles) y las filas utiles,
sin filas vacias, sin TOTAL y sin duplicados exactos. Todo lo descartado queda explicado en los avisos."""
import csv
import io
import re
from dataclasses import dataclass, field
from openpyxl import load_workbook
from rasgos_celdas import limpiar

MAX_FILAS_TITULO = 20                                  # la cabecera se busca en las primeras filas con algo
SEPARADORES = [";", ",", "\t"]
RE_TOTAL = re.compile(r"^(TOTAL(ES)?|SUMA|SUBTOTAL)\b")
RE_DATO = re.compile(r"@|\d[\d\s\-/.,:]{4,}\d")         # email, telefono o fecha: esto no es una cabecera


class ErrorLectura(Exception):
    """El archivo no se puede leer: el mensaje se ensena tal cual al gerente."""


@dataclass
class Tabla:
    hoja: str
    fila_cabecera: int                                  # numero de fila en el archivo (empieza en 1)
    cabecera: list
    filas: list                                         # listas de valores, tan largas como la cabecera
    num_filas: list                                     # numero de fila en el archivo de cada fila
    avisos: list = field(default_factory=list)


def leer(nombre, datos):
    ext = nombre.rsplit(".", 1)[-1].lower() if "." in nombre else ""
    if ext in ("csv", "txt"):
        hojas = {"CSV": leer_csv(datos)}
    elif ext in ("xlsx", "xlsm"):
        try:
            hojas = leer_xlsx(datos)
        except Exception as e:
            raise ErrorLectura("No se puede abrir el Excel. Comprueba que no esta danado ni protegido con contrasena.") from e
    elif ext == "xls":
        raise ErrorLectura("Es un Excel antiguo (.xls). Abrelo en Excel, guardalo como .xlsx y vuelve a subirlo.")
    else:
        raise ErrorLectura("Solo se pueden subir archivos .xlsx o .csv.")
    tablas, avisos = [], []
    for hoja, filas in hojas.items():
        t = tabla(hoja, filas)
        if t:
            tablas.append(t)
        else:
            avisos.append(f"Hoja '{hoja}': no tiene una tabla con cabecera, no se importa.")
    if not tablas:
        raise ErrorLectura("No se ha encontrado ninguna tabla con cabecera y datos.")
    return tablas, avisos


def leer_csv(datos):
    if datos.startswith(b"\xef\xbb\xbf"):
        texto = datos.decode("utf-8-sig")
    else:
        try:
            texto = datos.decode("utf-8")
        except UnicodeDecodeError:
            texto = datos.decode("cp1252", errors="replace")   # CSV de Excel espanol en Windows
    lineas = [l for l in texto.splitlines() if l.strip()][:30]
    # El separador es el que aparece el mismo numero de veces (y alguna) en casi todas las lineas
    sep = max(SEPARADORES, key=lambda s: (sorted(l.count(s) for l in lineas)[len(lineas) // 4] if lineas else 0))
    return [list(f) for f in csv.reader(io.StringIO(texto), delimiter=sep)]


def leer_xlsx(datos):
    wb = load_workbook(io.BytesIO(datos), data_only=True)
    hojas = {}
    for ws in wb.worksheets:
        filas = [list(f) for f in ws.iter_rows(min_row=1, min_col=1, max_row=ws.max_row, max_col=ws.max_column, values_only=True)]
        for rango in ws.merged_cells.ranges:           # celda combinada: el valor vale para todo el rango
            v = filas[rango.min_row - 1][rango.min_col - 1]
            for f in range(rango.min_row - 1, rango.max_row):
                for c in range(rango.min_col - 1, rango.max_col):
                    filas[f][c] = v
        hojas[ws.title] = filas
    return hojas


def valor(v):
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def es_cabecera(fila):
    llenas = [v for v in fila if v is not None]
    return len(llenas) >= 2 and all(isinstance(v, str) and not RE_DATO.search(v) for v in llenas)


def tabla(hoja, crudas):
    filas = [(i + 1, [valor(v) for v in f]) for i, f in enumerate(crudas)]
    filas = [(n, f) for n, f in filas if any(v is not None for v in f)]
    if not filas:
        return None
    ancho = max(sum(v is not None for v in f) for _, f in filas)
    pos = next((k for k, (_, f) in enumerate(filas[:MAX_FILAS_TITULO])
                if es_cabecera(f) and sum(v is not None for v in f) >= ancho / 2), None)
    if pos is None:
        return None
    n_cab, cab = filas[pos]
    avisos = [f"Filas 1-{n_cab - 1}: titulo antes de la cabecera, no se importa."] if n_cab > 1 else []
    # Cabecera doble: la de arriba agrupa (valores repetidos por celdas combinadas) y la de abajo detalla
    if pos + 1 < len(filas) and es_cabecera(filas[pos + 1][1]):
        abajo = filas[pos + 1][1]
        if len({v for v in cab if v}) < len({v for v in abajo if v}):
            cab = [f"{a} / {b}" if a and b and a != b else (b or a) for a, b in zip(cab, abajo + [None] * len(cab))]
            pos += 1
    ultima = max(c for c, v in enumerate(cab) if v is not None)
    cab = [v if v is not None else f"(columna {c + 1})" for c, v in enumerate(cab[:ultima + 1])]
    t = Tabla(hoja, n_cab, cab, [], [], avisos)
    vistas, cab_limpia = set(), [limpiar(v) for v in cab]
    for n, f in filas[pos + 1:]:
        f = (f + [None] * len(cab))[:len(cab)]
        if all(v is None for v in f):
            continue                                    # solo tenia algo fuera de la tabla
        primera = next(v for v in f if v is not None)
        clave = tuple(limpiar(v) if v is not None else "" for v in f)
        if isinstance(primera, str) and RE_TOTAL.match(limpiar(primera)):
            t.avisos.append(f"Fila {n}: es una fila de total, no se importa.")
        elif list(clave) == cab_limpia:
            t.avisos.append(f"Fila {n}: repite la cabecera, no se importa.")
        elif clave in vistas:
            t.avisos.append(f"Fila {n}: repetida exactamente, no se importa dos veces.")
        else:
            vistas.add(clave)
            t.filas.append(f)
            t.num_filas.append(n)
    return t if t.filas else None
